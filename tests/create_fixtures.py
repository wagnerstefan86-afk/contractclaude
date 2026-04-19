"""Generate test fixture files for parser verification."""
import os
import sys

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")


def create_pdf():
    """Create a 2-page PDF with contract text including audit clauses."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.units import cm

    path = os.path.join(FIXTURES_DIR, "test_contract.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4)
    styles = getSampleStyleSheet()

    heading1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=16, spaceAfter=12)
    heading2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceAfter=10)
    body = styles["BodyText"]

    story = [
        Paragraph("Rahmenvertrag IT-Services", heading1),
        Paragraph("1 Vertragsgegenstand", heading2),
        Paragraph(
            "Dieser Vertrag regelt die Erbringung von Managed Services "
            "zwischen dem Auftraggeber und dem Auftragnehmer. Der Vertrag "
            "umfasst den Betrieb der IT-Infrastruktur sowie zugehörige "
            "Support-Leistungen.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("2 Laufzeit und Kündigung", heading2),
        Paragraph(
            "Der Vertrag tritt am 01.01.2026 in Kraft und hat eine "
            "Mindestlaufzeit von 36 Monaten. Die Kündigungsfrist beträgt "
            "6 Monate zum Vertragsende.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("3 Service Level Agreement", heading2),
        Paragraph(
            "Der Auftragnehmer gewährleistet eine Verfügbarkeit von "
            "99.5% gemessen auf Monatsbasis. Geplante Wartungsfenster "
            "sind von der Messung ausgenommen.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("4 Haftung", heading2),
        Paragraph(
            "Die Haftung des Auftragnehmers ist auf den jährlichen "
            "Auftragswert begrenzt. Vertragsstrafen werden auf maximal "
            "10% des Monatsentgelts gedeckelt.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("5 Audit und Kontrolle", heading2),
        Paragraph(
            "Der Auftraggeber hat das Recht, jederzeit Audits und "
            "Prüfungen beim Auftragnehmer durchzuführen. Der Auftragnehmer "
            "stellt alle erforderlichen Unterlagen und Zugänge kostenfrei "
            "zur Verfügung. Drittprüfer können ohne vorherige Ankündigung "
            "eingesetzt werden.",
            body,
        ),
        Paragraph("5.1 Auditfrequenz", heading2),
        Paragraph(
            "Es gibt keine Begrenzung der Auditfrequenz. Inspections "
            "können unlimited durchgeführt werden, auch anytime outside "
            "business hours.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("6 Incident Management", heading2),
        Paragraph(
            "Sicherheitsvorfälle sind unverzüglich, spätestens innerhalb "
            "von 2 Stunden nach Feststellung zu melden. Der Auftragnehmer "
            "ist verpflichtet, bei jedem Security-Event eine Meldung "
            "an den Auftraggeber zu senden.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("7 Subunternehmer", heading2),
        Paragraph(
            "Der Einsatz von Subunternehmern bedarf der vorherigen "
            "schriftlichen Zustimmung des Auftraggebers. Die Anlage 3 "
            "enthält die genehmigten Unterauftragnehmer.",
            body,
        ),
        Spacer(1, 0.5 * cm),
        Paragraph("8 Datenschutz", heading2),
        Paragraph(
            "Die Parteien vereinbaren eine Auftragsverarbeitung gemäß "
            "Art. 28 DSGVO. Der Auftragnehmer verfügt über eine "
            "ISO 27001 Zertifizierung.",
            body,
        ),
    ]

    doc.build(story)
    print(f"Created: {path}")


def create_docx():
    """Create a DOCX with contract text using proper heading styles."""
    from docx import Document
    from docx.shared import Pt

    path = os.path.join(FIXTURES_DIR, "test_contract.docx")
    doc = Document()

    doc.add_heading("Rahmenvertrag IT-Services", level=1)

    doc.add_heading("1 Vertragsgegenstand", level=2)
    doc.add_paragraph(
        "Dieser Vertrag regelt die Erbringung von Managed Services "
        "zwischen dem Auftraggeber und dem Auftragnehmer."
    )

    doc.add_heading("2 Service Level Agreement", level=2)
    doc.add_paragraph(
        "Der Auftragnehmer gewährleistet eine Verfügbarkeit von "
        "99.5% gemessen auf Monatsbasis."
    )

    doc.add_heading("3 Audit und Kontrolle", level=2)
    doc.add_paragraph(
        "Der Auftraggeber hat das Recht, jederzeit Audits und "
        "Prüfungen beim Auftragnehmer durchzuführen. Der Auftragnehmer "
        "stellt alle erforderlichen Unterlagen kostenfrei zur Verfügung."
    )

    doc.add_heading("3.1 Auditfrequenz", level=3)
    doc.add_paragraph(
        "Es gibt keine Begrenzung der Auditfrequenz. Inspections "
        "können unlimited durchgeführt werden."
    )

    doc.add_heading("4 Incident Management", level=2)
    doc.add_paragraph(
        "Sicherheitsvorfälle sind unverzüglich zu melden. Die "
        "Meldepflicht gilt für jeden Vorfall ohne Schwellenwert."
    )

    doc.add_heading("5 Subunternehmer", level=2)
    doc.add_paragraph(
        "Der Einsatz von Subunternehmern bedarf der vorherigen "
        "Zustimmung. Gemäß Anlage 3 sind die genehmigten "
        "Unterauftragnehmer aufgeführt."
    )

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "SLA"
    table.cell(1, 0).text = "E-Mail"
    table.cell(1, 1).text = "99.9%"
    table.cell(2, 0).text = "Fileserver"
    table.cell(2, 1).text = "99.5%"

    doc.save(path)
    print(f"Created: {path}")


def create_xlsx():
    """Create an XLSX questionnaire with header row and data."""
    from openpyxl import Workbook

    path = os.path.join(FIXTURES_DIR, "test_questionnaire.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "InfoSec Fragebogen"

    headers = ["Nr.", "Anforderung", "Antwort", "Status", "Kommentar"]
    ws.append(headers)

    rows = [
        ["1", "Ist eine ISO 27001 Zertifizierung vorhanden?", "Ja", "Erfüllt", "Zertifikat liegt vor"],
        ["2", "Werden regelmäßige Penetrationstests durchgeführt?", "Ja, jährlich", "Erfüllt", ""],
        ["3", "Gibt es ein ISMS gemäß ISO 27001?", "Ja", "Erfüllt", "Scope: Managed Services"],
        ["4", "Wie werden Sicherheitsvorfälle gemeldet (Pflicht)?", "24h Erstmeldung", "Teilweise erfüllt", "Frist klären"],
        ["5", "Werden Audit-Rechte für den Auftraggeber eingeräumt?", "Ja, 1x jährlich", "Erfüllt", ""],
        ["6", "Gibt es ein BCM/ITSCM? (mandatory)", "Ja", "Erfüllt", "RTO: 4h"],
        ["7", "Werden Subunternehmer eingesetzt?", "Ja, siehe Anlage", "Klärungsbedarf", ""],
        ["8", "Welche Verfügbarkeit wird garantiert?", "99.5%", "Erfüllt", "SLA vorhanden"],
        ["9", "Gibt es Vertragsstrafen bei SLA-Verletzung?", "Ja, gedeckelt", "Erfüllt", "Max 10% Monatsentgelt"],
        ["10", "Ist eine Datenschutz-Folgenabschätzung vorhanden? (required)", "Nein", "Offen", "Muss erstellt werden"],
        ["11", "Werden Wartungsfenster explizit ausgenommen?", "Ja", "Erfüllt", ""],
        ["12", "Gibt es eine Exit-Regelung?", "Ja, §15", "Erfüllt", "Übergabefrist 6 Monate"],
    ]
    for row in rows:
        ws.append(row)

    wb.save(path)
    print(f"Created: {path}")


def main():
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    create_pdf()
    create_docx()
    create_xlsx()
    print("All fixtures created.")


if __name__ == "__main__":
    main()
