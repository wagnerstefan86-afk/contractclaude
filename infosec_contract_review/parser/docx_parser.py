"""DOCX parser using python-docx."""
from __future__ import annotations

import logging
import re

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from . import RawSegment

logger = logging.getLogger(__name__)

_HEADING_STYLE_PREFIX = "Heading"
_LIST_STYLE_PREFIXES = ("List", "list")


def parse_docx(file_path: str) -> list[RawSegment]:
    logger.info("Parsing DOCX: %s", file_path)
    try:
        doc = DocxDocument(file_path)
    except (PackageNotFoundError, Exception) as e:
        raise ValueError(f"Cannot open DOCX: {e}")

    segments: list[RawSegment] = []
    heading_stack: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else ""

        if style_name.startswith(_HEADING_STYLE_PREFIX):
            try:
                level = int(style_name.replace(_HEADING_STYLE_PREFIX, "").strip())
            except ValueError:
                level = 1
            heading_stack = heading_stack[:level - 1]
            heading_stack.append(text[:200])
            segments.append(RawSegment(
                text=text,
                heading_path=list(heading_stack),
                segment_type="heading",
                parse_quality="high",
            ))
        elif any(style_name.startswith(p) for p in _LIST_STYLE_PREFIXES):
            segments.append(RawSegment(
                text=text,
                heading_path=list(heading_stack),
                segment_type="list_item",
                parse_quality="high",
            ))
        else:
            segments.append(RawSegment(
                text=text,
                heading_path=list(heading_stack),
                segment_type="paragraph",
                parse_quality="high",
            ))

    for table in doc.tables:
        for row in table.rows:
            cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells_text:
                row_text = " | ".join(cells_text)
                segments.append(RawSegment(
                    text=row_text,
                    heading_path=list(heading_stack),
                    segment_type="table_cell",
                    parse_quality="medium",
                ))

    logger.info("DOCX parsed: %d segments from %s", len(segments), file_path)
    return segments
