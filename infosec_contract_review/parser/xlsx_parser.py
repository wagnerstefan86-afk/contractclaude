"""XLSX parser using openpyxl."""
from __future__ import annotations

import logging
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import RawSegment

logger = logging.getLogger(__name__)

_MANDATORY_PATTERNS = re.compile(
    r"\b(pflicht|mandatory|required|muss|obligatorisch)\b", re.IGNORECASE
)

_COLUMN_ROLE_HINTS = {
    "question": ["frage", "question", "anforderung", "requirement", "kriterium"],
    "answer": ["antwort", "answer", "response", "erfüllung"],
    "comment": ["kommentar", "comment", "bemerkung", "anmerkung", "hinweis"],
    "reference": ["referenz", "reference", "verweis", "nachweis", "dokument"],
    "status": ["status", "bewertung", "ergebnis", "result"],
}


def _guess_column_role(header: str) -> str:
    h = header.lower()
    for role, hints in _COLUMN_ROLE_HINTS.items():
        if any(hint in h for hint in hints):
            return role
    return "other"


def parse_xlsx(file_path: str) -> list[RawSegment]:
    logger.info("Parsing XLSX: %s", file_path)
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open XLSX: {e}")

    segments: list[RawSegment] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=False))
        if not rows:
            continue

        header_row = rows[0]
        headers = []
        col_roles = []
        for cell in header_row:
            val = str(cell.value).strip() if cell.value is not None else ""
            headers.append(val)
            col_roles.append(_guess_column_role(val))

        if not any(h for h in headers):
            continue

        for row_idx, row in enumerate(rows[1:], start=2):
            cells = []
            non_empty = False
            for cell in row:
                val = str(cell.value).strip() if cell.value is not None else ""
                cells.append(val)
                if val:
                    non_empty = True
            if not non_empty:
                continue

            parts = []
            for i, (header, val) in enumerate(zip(headers, cells)):
                if val:
                    parts.append(f"{header}: {val}" if header else val)

            text = " | ".join(parts)
            is_mandatory = bool(_MANDATORY_PATTERNS.search(text))

            first_col = get_column_letter(1)
            last_col = get_column_letter(len(cells))
            cell_range = f"{first_col}{row_idx}:{last_col}{row_idx}"

            segments.append(RawSegment(
                text=text,
                heading_path=[sheet_name],
                segment_type="questionnaire_item",
                parse_quality="high",
                extra={
                    "sheet_name": sheet_name,
                    "cell_range": cell_range,
                    "column_roles": col_roles[:len(cells)],
                    "is_mandatory": is_mandatory,
                },
            ))

    wb.close()
    logger.info("XLSX parsed: %d segments from %s", len(segments), file_path)
    return segments
